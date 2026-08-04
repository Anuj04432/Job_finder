"""
Step 1: Extract raw text from a resume file (PDF, DOCX, or image).

Now uses Microsoft's MarkItDown as the PRIMARY extractor for PDF and DOCX —
it preserves document structure (headings, lists, tables) as Markdown, which
tends to produce cleaner, better-structured text than raw text extraction
(and DOCX headings come through as real "#"/"##" markers, making section
detection much easier downstream).

Falls back automatically to the original pdfplumber/python-docx approach if:
  - MarkItDown's output looks too short (common for scanned/image-based PDFs,
    since MarkItDown's free path doesn't do OCR — that needs a paid LLM
    vision plugin, which we're deliberately not using here)
  - MarkItDown raises any error on a given file

Scanned/image PDFs and plain image files still go through local Tesseract
OCR (free, no API needed) — MarkItDown is not used for those.

Install dependencies:
    pip install markitdown[pdf,docx] pdfplumber PyMuPDF python-docx pytesseract pillow --break-system-packages

You also need the Tesseract binary installed on the system (not just the python package):
    Ubuntu/Debian: sudo apt-get install tesseract-ocr
    macOS:         brew install tesseract
    Windows:       https://github.com/UB-Mannheim/tesseract/wiki
"""

import os
import io
from pathlib import Path

from markitdown import MarkItDown
import pdfplumber
import fitz  # PyMuPDF
from docx import Document
import pytesseract
from PIL import Image

# --- Windows users: if Tesseract isn't on your PATH, uncomment and set this
# to wherever you installed it (default install location shown below):
# pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'


# Minimum characters we'd expect from a "real" text-based page.
# If extraction falls below this, we assume it's a scanned/image page and OCR it instead.
MIN_CHARS_PER_PAGE = 40

_markitdown_client = MarkItDown()


def _rough_page_count(file_path: str) -> int:
    """Quick page count via PyMuPDF, used to judge whether MarkItDown's
    output length is plausible for this document (scanned PDFs will look
    suspiciously short relative to their page count)."""
    try:
        doc = fitz.open(file_path)
        count = doc.page_count
        doc.close()
        return max(count, 1)
    except Exception:
        return 1


def extract_text_from_pdf_markitdown(file_path: str) -> str:
    """Primary PDF extraction path via MarkItDown."""
    result = _markitdown_client.convert(file_path)
    return (result.text_content or "").strip()


def extract_text_from_pdf_legacy(file_path: str) -> str:
    """
    Original fallback PDF extraction: pdfplumber for native text, with
    automatic per-page OCR fallback for scanned pages. Used only when
    MarkItDown's output looks too thin to trust.
    """
    text_chunks = []

    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages):
            page_text = page.extract_text() or ""

            if len(page_text.strip()) >= MIN_CHARS_PER_PAGE:
                text_chunks.append(page_text)
            else:
                # Fallback: render this page as an image and OCR it
                ocr_text = _ocr_pdf_page(file_path, page_num)
                text_chunks.append(ocr_text)

    return "\n\n".join(text_chunks).strip()


def extract_text_from_pdf(file_path: str) -> str:
    """
    Unified PDF extraction: tries MarkItDown first (better structure
    preservation). Falls back to the legacy pdfplumber+OCR approach if
    MarkItDown's output looks too short relative to the page count
    (a strong signal of a scanned/image-based PDF, which MarkItDown's
    free path can't OCR) or if MarkItDown raises an error.
    """
    page_count = _rough_page_count(file_path)

    try:
        text = extract_text_from_pdf_markitdown(file_path)
    except Exception:
        text = ""

    if len(text) >= MIN_CHARS_PER_PAGE * page_count:
        return text

    # MarkItDown's output looks too thin — likely a scanned PDF, or MarkItDown
    # failed silently on this file's structure. Fall back to the legacy path,
    # which has its own per-page OCR fallback built in.
    return extract_text_from_pdf_legacy(file_path)


def _ocr_pdf_page(file_path: str, page_num: int, zoom: int = 3) -> str:
    """Render a single PDF page to an image and run OCR on it."""
    doc = fitz.open(file_path)
    page = doc.load_page(page_num)

    # Higher zoom = higher resolution = better OCR accuracy
    matrix = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=matrix)
    img_bytes = pix.tobytes("png")

    image = Image.open(io.BytesIO(img_bytes))
    doc.close()

    return pytesseract.image_to_string(image)


def extract_text_from_docx_markitdown(file_path: str) -> str:
    """Primary DOCX extraction path via MarkItDown (preserves heading structure)."""
    result = _markitdown_client.convert(file_path)
    return (result.text_content or "").strip()


def extract_text_from_docx_legacy(file_path: str) -> str:
    """
    Original fallback DOCX extraction via python-docx. Used only if
    MarkItDown's output looks empty/too short for this file.
    """
    doc = Document(file_path)
    parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n".join(parts).strip()


def extract_text_from_docx(file_path: str) -> str:
    """
    Unified DOCX extraction: tries MarkItDown first (preserves heading
    levels as Markdown '#'/'##', which helps downstream section detection).
    Falls back to python-docx if MarkItDown's output is too short or errors.
    """
    try:
        text = extract_text_from_docx_markitdown(file_path)
    except Exception:
        text = ""

    if len(text) >= 20:
        return text

    return extract_text_from_docx_legacy(file_path)


def extract_text_from_image(file_path: str) -> str:
    """
    Extract text from a plain image file via local Tesseract OCR.
    (Deliberately NOT using MarkItDown here — its image handling relies on
    an LLM vision plugin that costs API money per call; local OCR is free.)
    """
    image = Image.open(file_path)
    return pytesseract.image_to_string(image).strip()


def extract_resume_text(file_path: str) -> str:
    """
    Unified entry point: detects file type by extension and routes to the
    right extractor. Returns raw extracted text.
    """
    ext = Path(file_path).suffix.lower()

    if not os.path.exists(file_path):
        raise FileNotFoundError(f"No such file: {file_path}")

    if ext == ".pdf":
        text = extract_text_from_pdf(file_path)
    elif ext in (".docx",):
        text = extract_text_from_docx(file_path)
    elif ext in (".png", ".jpg", ".jpeg", ".tiff", ".bmp"):
        text = extract_text_from_image(file_path)
    elif ext == ".doc":
        raise ValueError(
            "Legacy .doc format is not directly supported. "
            "Convert to .docx first (e.g. with LibreOffice: "
            "`libreoffice --headless --convert-to docx file.doc`)."
        )
    else:
        raise ValueError(f"Unsupported file type: {ext}")

    if not text or len(text.strip()) < 20:
        raise ValueError(
            "Extraction returned little or no text. The file may be corrupted, "
            "password-protected, or an unsupported scanned format."
        )

    return text


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("Usage: python extract_resume_text.py <path_to_resume>")
        sys.exit(1)

    path = sys.argv[1]
    extracted = extract_resume_text(path)
    print(extracted)